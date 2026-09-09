"""Run a job through the MCE Dryer HMB Calculator without touching the master.
usage: python3 run_case.py case.json out.xlsx   (case.json = {"Input label": value, ...} using the exact Inputs column-B labels)
Writes out.xlsx, recalculates it, and exports the Flow sheet to out_Flow.pdf."""
import sys, json, shutil, subprocess, os, glob
import openpyxl
HERE=os.path.dirname(os.path.abspath(__file__)); MASTER=os.path.join(HERE,'MCE_Dryer_HMB_Calculator.xlsx')
case=json.load(open(sys.argv[1])); out=sys.argv[2]
shutil.copy(MASTER,out)
wb=openpyxl.load_workbook(out); wi=wb['Inputs']
labels={wi.cell(r,2).value:r for r in range(1,wi.max_row+1) if wi.cell(r,2).value}
for k,v in case.items():
    if k.startswith('fuel:'):   # fuel table override: "fuel:Biogas 55% CH4" -> [HHV, lb/ft3, xs, stoich, h2o, co2]
        name=k[5:]; row=[r for r in range(13,19) if wi.cell(r,7).value==name][0]
        for j,val in enumerate(v): wi.cell(row,8+j).value=val
        continue
    if k not in labels: raise SystemExit(f'no input labelled {k!r}')
    wi.cell(labels[k],3).value=v
wb.save(out)
recalc=glob.glob('/root/.claude/skills/synced/*/xlsx/scripts/recalc.py')[0]
env=dict(os.environ,HOME=os.environ.get('LO_HOME',os.environ['HOME']))
print(subprocess.run(['python3',recalc,out,'240'],capture_output=True,text=True,env=env).stdout)
outdir=os.path.dirname(os.path.abspath(out)) or '.'
subprocess.run(['soffice','-env:UserInstallation=file://'+env['HOME']+'/loprof','--headless','--convert-to','pdf','--outdir',outdir,out],capture_output=True,env=env,timeout=300)
pdf=out[:-5]+'.pdf'
page=None
for p in range(1,80):
    t=subprocess.run(['pdftotext','-f',str(p),'-l',str(p),pdf,'-'],capture_output=True,text=True).stdout
    if 'DEHYDRATION' in t: page=p; break
subprocess.run(['pdfseparate','-f',str(page),'-l',str(page),pdf,out[:-5]+'_Flow.pdf']); os.remove(pdf)
wb=openpyxl.load_workbook(out,data_only=True)
for sh in ('Summary',):
    ws=wb[sh]
    for r in range(1,ws.max_row+1):
        lab=ws.cell(r,2).value; v=ws.cell(r,3).value
        if lab and v is not None and not str(lab).isupper(): print(f'{str(lab)[:50]:50s} {v}')
